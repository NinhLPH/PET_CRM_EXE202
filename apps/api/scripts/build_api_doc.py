"""Build the frontend-facing PetCare CRM v1 API reference."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = Path(__file__).resolve().parents[1] / "API_DOC.docx"

# Method, path, purpose, request, response. All paths are relative to the API origin.
GROUPS = [
    ("Xác thực và hồ sơ", [
        ("POST", "/auth/register", "Tạo tài khoản CUSTOMER; liên kết hồ sơ CRM có cùng SĐT nếu duy nhất.", "Body Register; không cần cookie.", "201 UserBrief {id, phone, role}; chưa tạo phiên."),
        ("POST", "/auth/login", "Đăng nhập CUSTOMER hoặc ADMIN.", "Body Login; không cần cookie.", "201 AuthLogin {role, userId}; Set-Cookie petcare_session."),
        ("POST", "/auth/logout", "Thu hồi phiên hiện tại.", "Cookie phiên; body rỗng.", "201 Success {success:true}; xóa cookie."),
        ("GET", "/me/profile", "Đọc hồ sơ của chủ nuôi đang đăng nhập.", "Cookie CUSTOMER.", "200 Profile {id, fullName, phone, address}."),
        ("PATCH", "/me/profile", "Sửa tên hoặc địa chỉ; SĐT chỉ đọc.", "Cookie CUSTOMER; body ProfilePatch.", "200 Profile."),
    ]),
    ("Pet của chủ nuôi", [
        ("GET", "/pets/mine", "Liệt kê pet thuộc tài khoản hiện tại.", "Cookie CUSTOMER; query page?, limit?.", "200 Page<Pet>."),
        ("GET", "/pets/{id}", "Xem một pet của mình.", "Cookie CUSTOMER; path id.", "200 Pet."),
        ("POST", "/pets", "Thêm pet cho mình.", "Cookie CUSTOMER; body PetCreate.", "201 Pet."),
        ("PATCH", "/pets/{id}", "Cập nhật pet của mình.", "Cookie CUSTOMER; path id; body PetPatch.", "200 Pet."),
        ("DELETE", "/pets/{id}", "Xóa pet chưa có booking hoặc reminder.", "Cookie CUSTOMER; path id.", "200 Success."),
    ]),
    ("Dịch vụ công khai và booking của chủ nuôi", [
        ("GET", "/services", "Danh sách dịch vụ đang bán.", "Công khai; query status=ACTIVE?, page?, limit?.", "200 Page<Service>."),
        ("GET", "/services/{id}/quote", "Lấy giá theo loài và cân nặng; dùng giá này khi đặt lịch.", "Công khai; path id; query species=DOG|CAT, weight=chuỗi thập phân.", "200 Quote {serviceId, species, weight, servicePriceId, basePrice}."),
        ("POST", "/bookings", "Đặt một pet và một dịch vụ tại thời điểm tương lai.", "Cookie CUSTOMER; header Idempotency-Key; body BookingCreate.", "201 Booking kèm services[1] chứa snapshot giá; request lặp cùng key có thể trả 201 Booking chi tiết."),
        ("GET", "/bookings/mine", "Lịch đặt của chủ nuôi, mới nhất trước.", "Cookie CUSTOMER; query page?, limit?.", "200 Page<Booking> kèm pet và services.service."),
        ("GET", "/bookings/{id}", "Xem chi tiết booking của mình.", "Cookie CUSTOMER; path id.", "200 BookingDetail kèm customer, pet, services, reminders; thông tin chốt giá chỉ hiển thị khi COMPLETED."),
    ]),
    ("Dashboard và khách hàng CRM", [
        ("GET", "/admin/dashboard", "Đếm booking chờ xác nhận, booking đã xác nhận hôm nay và reminder đến hạn.", "Cookie ADMIN.", "200 {pending, confirmedToday, remindersDue}."),
        ("GET", "/admin/customers", "Tìm theo tên hoặc SĐT, mới nhất trước.", "Cookie ADMIN; query q?, page?, limit?.", "200 Page<Customer>."),
        ("GET", "/admin/customers/{id}", "Hồ sơ 360 độ của khách.", "Cookie ADMIN; path id.", "200 Customer kèm pets, bookings (20 mới nhất, mỗi booking có services), activities (20 mới nhất)."),
        ("POST", "/admin/customers", "Tạo hồ sơ CRM chưa cần tài khoản.", "Cookie ADMIN; body CustomerCreate.", "201 Customer."),
        ("PATCH", "/admin/customers/{id}", "Sửa hồ sơ CRM; hồ sơ đã gắn tài khoản không đổi SĐT.", "Cookie ADMIN; path id; body CustomerPatch.", "200 Customer."),
        ("DELETE", "/admin/customers/{id}", "Xóa khách chưa có tài khoản hoặc dữ liệu phụ thuộc.", "Cookie ADMIN; path id.", "200 Success."),
        ("GET", "/admin/customers/{id}/pets", "Liệt kê pet của một khách.", "Cookie ADMIN; path id; query page?, limit?.", "200 Page<Pet>."),
        ("POST", "/admin/customers/{id}/pets", "Thêm pet vào hồ sơ khách.", "Cookie ADMIN; path id; body PetCreate.", "201 Pet."),
        ("PATCH", "/admin/pets/{petId}", "Sửa pet trong CRM.", "Cookie ADMIN; path petId; body PetPatch.", "200 Pet."),
        ("DELETE", "/admin/pets/{petId}", "Xóa pet chưa có booking hoặc reminder.", "Cookie ADMIN; path petId.", "200 Success."),
    ]),
    ("Danh mục dịch vụ và bảng giá", [
        ("GET", "/admin/services", "Liệt kê cả dịch vụ ACTIVE và INACTIVE.", "Cookie ADMIN; query page?, limit?.", "200 Page<Service>."),
        ("POST", "/admin/services", "Thêm dịch vụ; dịch vụ mới ở trạng thái INACTIVE.", "Cookie ADMIN; body ServiceCreate.", "201 Service."),
        ("PATCH", "/admin/services/{id}", "Sửa tên, mô tả hoặc thời lượng.", "Cookie ADMIN; path id; body ServicePatch.", "200 Service."),
        ("PATCH", "/admin/services/{id}/status", "Mở hoặc ngừng bán; mở bán cần ít nhất một khoảng giá ACTIVE.", "Cookie ADMIN; path id; body StatusBody.", "200 Service."),
        ("GET", "/admin/services/{id}/prices", "Liệt kê khoảng giá ACTIVE của dịch vụ.", "Cookie ADMIN; path id; query species=DOG|CAT?, page?, limit?.", "200 Page<ServicePrice>; không trả giá INACTIVE."),
        ("POST", "/admin/services/{id}/prices", "Thêm khoảng giá cho DOG hoặc CAT.", "Cookie ADMIN; path id; body PriceBody.", "201 ServicePrice."),
        ("PATCH", "/admin/services/{id}/prices/{priceId}", "Sửa toàn bộ loài, khoảng cân nặng và giá của một rule.", "Cookie ADMIN; path id, priceId; body PriceBody (đầy đủ trường bắt buộc).", "200 ServicePrice."),
        ("PATCH", "/admin/prices/{priceId}/status", "Bật hoặc tắt khoảng giá; bật không được chồng khoảng.", "Cookie ADMIN; path priceId; body StatusBody.", "200 ServicePrice."),
        ("GET", "/admin/services/{id}/reminder-config", "Lấy cấu hình nhắc ACTIVE của dịch vụ.", "Cookie ADMIN; path id.", "200 ReminderConfig hoặc null nếu chưa cấu hình."),
        ("PUT", "/admin/services/{id}/reminder-config", "Tạo hoặc cập nhật chu kỳ nhắc.", "Cookie ADMIN; path id; body {reminderDays: 8..364}.", "200 ReminderConfig."),
    ]),
    ("Booking cho admin", [
        ("GET", "/admin/bookings", "Lọc booking theo trạng thái và ngày hẹn Việt Nam.", "Cookie ADMIN; query status?, date=YYYY-MM-DD?, page?, limit?.", "200 Page<Booking> kèm customer, pet, services.service."),
        ("GET", "/admin/bookings/{id}", "Xem đầy đủ booking, phụ phí và reminder.", "Cookie ADMIN; path id.", "200 BookingDetail."),
        ("POST", "/admin/bookings/{id}/confirm", "Chuyển PENDING thành CONFIRMED.", "Cookie ADMIN; path id; body rỗng.", "201 Booking."),
        ("POST", "/admin/bookings/{id}/cancel", "Hủy PENDING hoặc CONFIRMED.", "Cookie ADMIN; path id; body CancelBody? hoặc {}.", "201 Booking có status=CANCELLED, cancelledAt, cancellationReason."),
        ("POST", "/admin/bookings/{id}/complete", "Chốt giá và tạo reminder trong một transaction.", "Cookie ADMIN; path id; body CompleteBody.", "201 BookingDetail có services, surcharges, reminders và finalTotal."),
    ]),
    ("Reminder và hoạt động CRM", [
        ("GET", "/admin/reminders", "Liệt kê reminder, ngày cũ trước; due=true chỉ lấy chưa xử lý đến hết hôm nay.", "Cookie ADMIN; query due=true|false?, page?, limit?.", "200 Page<Reminder> kèm customer, pet, service."),
        ("GET", "/admin/reminders/{id}", "Xem bối cảnh trước khi liên hệ.", "Cookie ADMIN; path id.", "200 ReminderDetail kèm customer, pet.activities, service, booking.services, booking.surcharges."),
        ("POST", "/admin/reminders/{id}/contact", "Đánh dấu đã liên hệ và tạo activity.", "Cookie ADMIN; path id; body ContactBody.", "201 Reminder có status=CONTACTED, contactedAt, contactMethod, note."),
        ("GET", "/admin/customers/{id}/activities", "Lịch sử hoạt động CRM của khách.", "Cookie ADMIN; path id; query page?, limit?.", "200 Page<Activity> kèm user {id, phone}."),
        ("POST", "/admin/customers/{id}/activities", "Tạo ghi chú thủ công cho khách hoặc pet của khách.", "Cookie ADMIN; path id; body ActivityBody.", "201 Activity có type=NOTE."),
    ]),
]

SCHEMAS = [
    ("Register", "fullName: string 1..100; phone: 10 chữ số bắt đầu bằng 0; password: string >=8; confirmPassword: giống password."),
    ("Login", "phone: 10 chữ số bắt đầu bằng 0; password: string."),
    ("ProfilePatch", "fullName?: string 1..100; address?: string <=255. Không gửi phone."),
    ("PetCreate", "name: string 1..100; species: DOG|CAT; weight: string >0 đến 999.99, tối đa 2 số lẻ; breed?: string <=100; allergyNote?: string; specialNote?: string."),
    ("PetPatch", "Các trường của PetCreate đều tùy chọn; nếu gửi weight phải hợp lệ."),
    ("CustomerCreate", "fullName: string 1..100; phone: 10 chữ số bắt đầu bằng 0; address?: string <=255; note?: string."),
    ("CustomerPatch", "Các trường của CustomerCreate đều tùy chọn; SĐT không được trùng và không đổi khi hồ sơ đã gắn tài khoản."),
    ("ServiceCreate", "serviceName: string 1..100; description?: string; estimatedDuration?: integer >=1 (phút)."),
    ("ServicePatch", "Các trường của ServiceCreate đều tùy chọn."),
    ("StatusBody", "status: ACTIVE|INACTIVE."),
    ("PriceBody", "species: DOG|CAT; minWeight: string >=0; maxWeight?: string > minWeight (bỏ qua nghĩa là không giới hạn); price: integer VND 1..9,999,999,999."),
    ("BookingCreate", "petId: ID string; serviceId: ID string; bookingDate: ISO 8601 có múi giờ, ở tương lai; expectedBasePrice: integer VND 1..9,999,999,999; note?: string."),
    ("CancelBody", "reason?: string <=255."),
    ("CompleteBody", "surcharges: array tối đa 50 phần tử (bắt buộc, có thể []), mỗi phần tử {name: string 1..100, amount: integer VND 1..9,999,999,999, note?: string <=255}; discount: integer VND >=0."),
    ("ContactBody", "method: PHONE|ZALO|OTHER; note?: string."),
    ("ActivityBody", "type: NOTE; content: string không rỗng; petId?: ID string của pet thuộc khách."),
]

RESPONSES = [
    ("Page<T>", "{items: T[], total: number, page: number, limit: number}; page mặc định 1, limit mặc định 20, tối đa 100."),
    ("Profile", "{id, fullName, phone, address}; address có thể null."),
    ("Pet", "{id, customerId, name, species, breed, weight, allergyNote, specialNote, createdAt, updatedAt}."),
    ("Customer", "{id, userId, fullName, phone, address, note, createdAt, updatedAt}; trường tùy chọn có thể null."),
    ("Service", "{id, serviceName, description, estimatedDuration, status, createdAt, updatedAt}."),
    ("ServicePrice", "{id, serviceId, species, minWeight, maxWeight, price, effectiveFrom, effectiveTo, status, createdAt, updatedAt}; maxWeight=null là vô hạn."),
    ("ReminderConfig", "{id, serviceId, reminderDays, status, createdAt, updatedAt}."),
    ("Booking", "{id, customerId, petId, bookingDate, status, note, estimatedTotal, finalTotal, discountAmount, completedAt, cancelledAt, cancellationReason, createdAt, updatedAt}; tùy route có thêm pet, customer, services, surcharges, reminders."),
    ("BookingService", "{id, bookingId, serviceId, servicePriceId, petWeightSnapshot, basePrice, finalPrice, createdAt}; có thể thêm service trong danh sách."),
    ("Surcharge", "{id, bookingId, bookingServiceId, surchargeName, amount, note, createdAt}."),
    ("Reminder", "{id, petId, customerId, bookingId, serviceId, completedDate, reminderDate, status, contactedAt, contactMethod, note, createdAt}."),
    ("Activity", "{id, customerId, petId, reminderId, type, content, createdBy, createdAt}."),
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table):
    tbl_pr = table._tbl.tblPr
    item = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        child = OxmlElement(f"w:{edge}")
        child.set(qn("w:val"), "single")
        child.set(qn("w:sz"), "4")
        child.set(qn("w:color"), "D9D9D9")
        item.append(child)
    tbl_pr.append(item)


def write_cell(cell, value, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.style = "Normal"
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(value)
    run.bold = bold
    run.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, heads, rows, widths):
    result = doc.add_table(rows=1, cols=len(heads))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    result._tbl.tblPr.append(layout)
    for i, head in enumerate(heads):
        result.columns[i].width = Inches(widths[i])
        result.rows[0].cells[i].width = Inches(widths[i])
        write_cell(result.rows[0].cells[i], head, True)
        shade(result.rows[0].cells[i], "EDEDED")
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    result.rows[0]._tr.get_or_add_trPr().append(header)
    for row in rows:
        new_row = result.add_row()
        cells = new_row.cells
        cant_split = OxmlElement("w:cantSplit")
        new_row._tr.get_or_add_trPr().append(cant_split)
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            write_cell(cells[i], value)
    borders(result)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_para(doc, text, style=None):
    p = doc.add_paragraph(text, style)
    p.paragraph_format.space_after = Pt(5)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    sec.left_margin = sec.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    for name, size in (("Title", 20), ("Heading 1", 13), ("Heading 2", 10)):
        s = styles[name]
        s.font.name = "Aptos"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(11)
        s.paragraph_format.space_after = Pt(5)
    title_style_pr = styles["Title"]._element.get_or_add_pPr()
    for border in title_style_pr.findall(qn("w:pBdr")):
        title_style_pr.remove(border)
    doc.add_paragraph("PetCare CRM API v1 cho Frontend", "Title")
    add_para(doc, "Tài liệu mô tả 45 endpoint đang có trong backend, cách gửi request, dạng response và các quy tắc FE cần xử lý. Đường dẫn trong bảng được ghép trực tiếp với API origin (không có tiền tố /api/v1). Swagger UI tại /openapi khi server chạy.")
    doc.add_heading("Quy ước tích hợp", 1)
    for line in [
        "Body ghi dữ liệu dùng JSON và Content-Type: application/json. Các route CUSTOMER và ADMIN dùng cookie petcare_session; với FE khác origin hãy gửi credentials: 'include' và cấu hình origin trong CORS_ORIGINS. Cookie HttpOnly không được đọc bằng JavaScript.",
        "ID BIGINT là chuỗi số trong request và response. Tiền là số nguyên VND; cân nặng là chuỗi thập phân tối đa hai chữ số lẻ. Thời điểm trả về theo ISO 8601 UTC; query date và ngày reminder xét theo Asia/Ho_Chi_Minh.",
        "GET danh sách có page và limit (mặc định 1 và 20, limit tối đa 100). Path {id}, {petId}, {priceId} đều là ID chuỗi dương. POST thường trả HTTP 201; PATCH, PUT, DELETE và GET trả 200 theo cấu hình hiện tại.",
        "Trình tự đặt lịch: lấy GET /services/{id}/quote, gửi basePrice vào expectedBasePrice của POST /bookings cùng một Idempotency-Key duy nhất cho thao tác. Nếu gặp PRICE_CHANGED, hiển thị giá mới và hỏi người dùng xác nhận lại."
    ]:
        add_para(doc, line)
    doc.add_heading("Lỗi chung", 1)
    add_para(doc, 'Lỗi có dạng {"code":"HTTP_400","message":"..."}. Validation trả code=VALIDATION_ERROR và errors:[{field,messages:string[]}]. 401 là thiếu hoặc hết phiên; 403 là sai vai trò, quyền sở hữu hoặc Origin; 404 không tìm thấy; 409 là trùng dữ liệu, giá, hoặc xung đột trạng thái; 429 là quá nhiều lần đăng nhập/đăng ký. Trường hợp đổi giá trả {"code":"PRICE_CHANGED","message":"Giá đã thay đổi","currentQuote":{"basePrice":250000,"servicePriceId":"7"}}.')
    doc.add_heading("Request body", 1)
    table(doc, ["Tên", "Trường và ràng buộc"], SCHEMAS, [1.25, 5.7])
    doc.add_heading("Response model", 1)
    add_para(doc, "Các model Prisma được trả về theo tên trường camelCase. ID, ngày giờ và Decimal được chuyển kiểu theo quy ước trên; trường có quan hệ chỉ xuất hiện khi route mô tả có kèm quan hệ đó. Giá trị nullable có thể là null.")
    table(doc, ["Tên", "Trường chính"], RESPONSES, [1.25, 5.7])
    count = 0
    for group, endpoints in GROUPS:
        doc.add_heading(group, 1)
        rows = []
        for method, path, purpose, request, response in endpoints:
            rows.append((f"{method} {path}", purpose, request, response))
            count += 1
        table(doc, ["Endpoint", "Mô tả", "Request", "Response"], rows, [1.55, 1.6, 1.85, 1.95])
    assert count == 45, count
    doc.add_heading("Ví dụ tích hợp", 1)
    examples = [
        ("Đăng nhập", 'POST /auth/login\n{"phone":"0901234567","password":"secret123"}\n201 {"role":"CUSTOMER","userId":"12"} + Set-Cookie'),
        ("Báo giá", 'GET /services/3/quote?species=DOG&weight=7.00\n200 {"serviceId":"3","species":"DOG","weight":"7","servicePriceId":"7","basePrice":250000}'),
        ("Đặt lịch", 'POST /bookings  |  Idempotency-Key: booking-form-uuid\n{"petId":"4","serviceId":"3","bookingDate":"2026-11-01T09:00:00+07:00","expectedBasePrice":250000,"note":"Tắm nhẹ"}\n201 Booking có status=PENDING, estimatedTotal=250000 và services[0].basePrice=250000'),
        ("Hoàn thành", 'POST /admin/bookings/9/complete\n{"surcharges":[{"name":"Gỡ rối lông","amount":50000}],"discount":20000}\n201 Booking có status=COMPLETED, finalTotal=280000 và reminders mới.'),
    ]
    for title, content in examples:
        doc.add_heading(title, 2)
        p = add_para(doc, content)
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(8)
    doc.add_heading("Lưu ý về trạng thái", 1)
    add_para(doc, "Booking: PENDING → CONFIRMED → COMPLETED; ADMIN có thể hủy từ PENDING hoặc CONFIRMED. Hoàn thành cần cấu hình reminder ACTIVE; finalTotal = basePrice snapshot + tổng phụ phí - discount và không được âm. Reminder liên hệ thành công chuyển CONTACTED và biến khỏi danh sách due=true. GET /admin/reminders không lọc trạng thái khi thiếu due=true (kể cả due=false).")
    doc.core_properties.title = "PetCare CRM API v1 cho Frontend"
    doc.core_properties.subject = "Tài liệu request response cho 45 endpoint backend"
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({count} endpoints)")


if __name__ == "__main__":
    main()
